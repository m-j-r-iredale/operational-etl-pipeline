# Operational ETL Pipeline - Audit Workflow
# Portfolio version based on the revised project article.

import os
import logging

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from sqlalchemy import URL, create_engine, text

logging.basicConfig(format="%(levelname)s:%(message)s", level=logging.INFO)

raw_data = r"C:\RedactedL\fake_audit_logging_spreadsheet.xlsm"
return_wb = r"C:\Redacted\audit_return.xlsx"

bank_holidays = np.array([
    "2026-01-01",
    "2026-04-03",
    "2026-04-06",
    "2026-05-04",
    "2026-05-25",
    "2026-08-31",
    "2026-12-25",
    "2026-12-28",
], dtype="datetime64[D]")

def extract(raw_data, sheet_name=0):
    return pd.read_excel(raw_data, sheet_name=sheet_name)

logging.info(" File extraction starting.")
try:
    df = extract(raw_data)
    logging.info(" File extracted successfully.")
    print(f"The data has {df.shape[0]} rows and {df.shape[1]} columns.")
    int_shape = df.shape
except Exception:
    logging.error(" Could not extract file!")
    raise

logging.info(" Standard data cleaning started.")

df.dropna(subset=df.columns.drop("UUID"), how="all", inplace=True)
df.rename(columns={"UUID": "transaction_id"}, inplace=True)
print(
    f"The dataframe went from {int_shape[0]} rows to {df.shape[0]} "
    f"and {int_shape[1]} columns to {df.shape[1]}."
)

df_unworked = df[df["Status"].isna() & df["Allocated to"].isna()]
df = df[df["Status"].notna()]
logging.info(" Unworked cases split from ongoing cases.")

df["Account Name"] = df["Account Name"].str.title().str.strip()
df["Description of Request (keep it short)"] = (
    df["Description of Request (keep it short)"].str.title().str.strip()
)
df["If 'Returned', why?"] = df["If 'Returned', why?"].str.lower().str.strip()
logging.info(" Text columns cleaned.")

df["Date Received"] = pd.to_datetime(
    df["Date Received"], format="%d/%m/%Y", errors="coerce"
)
df["Date Process Started"] = pd.to_datetime(
    df["Date Process Started"], format="%d/%m/%Y", errors="coerce"
)
df["Date Process Completed"] = pd.to_datetime(
    df["Date Process Completed"], format="%d/%m/%Y", errors="coerce"
)
logging.info(" Date columns standardised.")

print(f"Date Received has {df['Date Received'].isna().sum()} null values.")
print(
    f"Date Process Started has "
    f"{df['Date Process Started'].isna().sum()} null values."
)
print(
    f"Date Process Completed has "
    f"{df['Date Process Completed'].isna().sum()} null values."
)

condition = (
    df["Date Received"].notna()
    & df["Date Process Completed"].notna()
    & df["Total Days to Complete"].isna()
)

print(
    f"Total Days to Complete had "
    f"{df['Total Days to Complete'].isna().sum()} null values."
)
df.loc[condition, "Total Days to Complete"] = np.busday_count(
    df.loc[condition, "Date Received"].values.astype("datetime64[D]"),
    df.loc[condition, "Date Process Completed"].values.astype("datetime64[D]"),
    holidays=bank_holidays,
)
print(
    f"Total Days to Complete now has "
    f"{df['Total Days to Complete'].isna().sum()} null values."
)
logging.info(" Total Days to Complete column verified.\n")

logging.info(" Error logging process starting.")
df["errors"] = ""

df.loc[df["Allocated to"].isna(), "errors"] += "missing_allocation_initials; "
df.loc[df["Date Process Started"].isna(), "errors"] += "missing_start_date; "
df.loc[
    df["Date Process Completed"].isna()
    & ((df["Status"] == "Completed") | (df["Status"] == "Returned")),
    "errors",
] += "missing_completed_date; "
df.loc[df["Account Number"].isna(), "errors"] += "missing_account_number(s); "
df.loc[df["Account Name"].isna(), "errors"] += "missing_account_name; "
df.loc[
    df["Type of Audit (Paper or Portal)"].isna(), "errors"
] += "missing_audit_type; "
df.loc[
    df["Audit or Change of Auditor"].isna(), "errors"
] += "audit_or_change_missing; "
df.loc[df["Date Received"].isna(), "errors"] += "missing_date_received; "
df.loc[
    df["Date Received"] > df["Date Process Started"], "errors"
] += "date_received_>_date_process_srt; "
df.loc[
    df["Date Received"] > df["Date Process Completed"], "errors"
] += "date_received_>_date_process_cmp; "
df.loc[
    df["Date Process Started"] > df["Date Process Completed"], "errors"
] += "date_process_srt_>_date_process_cmp; "
logging.info(" All errors logged.")

logging.info(" Upload dataframe creation starting.")
load_condition = (df["errors"] == "") & df["Status"].notna()
df_load = df[load_condition]

df_load["Type of Audit (Paper or Portal)"] = (
    "Audit " + df_load["Type of Audit (Paper or Portal)"]
)
logging.info(" Data ready for upload.")
logging.info(" Create Return sheet.")

df_incomplete = df_load[~df_load["Status"].isin(["Completed", "Returned"])]
df_errors = df[df["errors"] != ""]
df_return = pd.concat([df_errors, df_incomplete, df_unworked])
logging.info(" Return sheet created.")

logging.info(" Preparing Return File for delivery.")
audit_workbook = load_workbook(return_wb)
audit_sheet = audit_workbook["Audit Log"]

for row in audit_sheet.iter_rows(min_row=2, max_row=audit_sheet.max_row):
    for cell in row:
        cell.value = None

if df_return.empty:
    print("Return data is empty")
else:
    for row_index, row in enumerate(df_return.itertuples(index=False), start=2):
        for column_index, value in enumerate(row, start=1):
            audit_sheet.cell(
                row=row_index,
                column=column_index,
                value=value,
            )
    audit_workbook.save(return_wb)
    logging.info(" Return File sent successfully")

logging.info(" Audit-specific transformation sequence started.")
df_load["agent_role"] = "Allocated"
logging.info(" Agent role added.")

df_load_bcp = df_load.copy()
logging.info(" Backup created.")

df_load["Date Received"] = (
    df_load["Date Received"].dt.strftime("%Y%m%d").astype("Int64")
)
df_load["Date Process Started"] = (
    df_load["Date Process Started"].dt.strftime("%Y%m%d").astype("Int64")
)
df_load["Date Process Completed"] = (
    df_load["Date Process Completed"].dt.strftime("%Y%m%d").astype("Int64")
)
logging.info(" Date keys created.")

db_host = os.getenv("DB_HOST")
db_port = os.getenv("DB_PORT")
db_name = os.getenv("DB_NAME")
db_user = os.getenv("DB_USER")
db_password = os.getenv("DB_PASSWORD")

db_url = URL.create(
    drivername="postgresql+psycopg",
    username=db_user,
    password=db_password,
    host=db_host,
    port=db_port,
    database=db_name,
)

engine = create_engine(db_url, pool_pre_ping=True)
logging.info(" Database connection ready.")

logging.info(" Create mapping function.")

def map_existing_dim(
    data,
    engine,
    sheet_col_name,
    db_table_name,
    db_col_name,
    db_key_col,
):
    sql_query = text(f"""
        SELECT
            {db_col_name},
            {db_key_col}
        FROM
            {db_table_name}
    """)

    df_db = pd.read_sql(sql_query, engine)
    mapping_df = df_db.set_index(db_col_name)[db_key_col]
    data[db_key_col] = data[sheet_col_name].map(mapping_df)
    return data

logging.info(" Call mapping functions.")

df_load = map_existing_dim(
    data=df_load,
    engine=engine,
    sheet_col_name="Allocated to",
    db_table_name="dim_agent",
    db_col_name="agent_id",
    db_key_col="agent_key",
)
logging.info(" Agent keys mapped.")

df_load = map_existing_dim(
    data=df_load,
    engine=engine,
    sheet_col_name="Type of Audit (Paper or Portal)",
    db_table_name="dim_task_type",
    db_col_name="task_type",
    db_key_col="task_key",
)
logging.info(" Task type mapped.")

logging.info(" Beginning Admin Fact table SQL call, upload and key exchange.")

fact_adm_sql = text("""
INSERT INTO fact_admin(
    date_key,
    date_start_key,
    date_complete_key,
    task_key,
    task_status,
    days_to_complete,
    source_organisation_name,
    transaction_id)

VALUES(
    :date_key,
    :date_start_key,
    :date_complete_key,
    :task_key,
    :task_status,
    :days_to_complete,
    :source_organisation_name,
    :transaction_id)

ON CONFLICT (transaction_id)
    DO UPDATE SET
        date_key = EXCLUDED.date_key,
        date_start_key = EXCLUDED.date_start_key,
        date_complete_key = EXCLUDED.date_complete_key,
        task_key = EXCLUDED.task_key,
        task_status = EXCLUDED.task_status,
        days_to_complete = EXCLUDED.days_to_complete,
        source_organisation_name = EXCLUDED.source_organisation_name,
        updated_at = CURRENT_TIMESTAMP
    RETURNING transaction_key
""")

df_load["transaction_key"] = pd.Series(pd.NA, index=df_load.index, dtype="Int64")
logging.info(" Transaction Key column created.")
logging.info(" Transaction Key exchange loop initiated.")

with engine.begin() as conn:
    for index, row in df_load.iterrows():
        transaction_key = conn.execute(
            fact_adm_sql,
            {
                "transaction_id": row["transaction_id"],
                "date_key": int(row["Date Received"]),
                "date_start_key": int(row["Date Process Started"]),
                "date_complete_key": (
                    None
                    if pd.isna(row["Date Process Completed"])
                    else int(row["Date Process Completed"])
                ),
                "task_key": int(row["task_key"]),
                "task_status": row["Status"],
                "days_to_complete": (
                    None
                    if pd.isna(row["Total Days to Complete"])
                    else int(row["Total Days to Complete"])
                ),
                "source_organisation_name": row["Account Name"],
            },
        ).scalar_one()

        df_load.loc[index, "transaction_key"] = transaction_key

logging.info(" Fact Table rows uploaded.")

logging.info(" Creating Agent Bridge table.")
bridge_agt_df = df_load[["transaction_key", "agent_key", "agent_role"]]

bridge_agent_sql = text("""
    INSERT INTO bridge_admin_agent(
        transaction_key,
        agent_key,
        agent_role
    )
    VALUES(
        :transaction_key,
        :agent_key,
        :agent_role
    )
    ON CONFLICT (
        transaction_key,
        agent_role
    )
    DO UPDATE SET
        agent_key = EXCLUDED.agent_key,
        updated_at = CURRENT_TIMESTAMP
    WHERE bridge_admin_agent.agent_key
        IS DISTINCT FROM EXCLUDED.agent_key
""")

logging.info(" Making Agent Bridge SQL call.")
with engine.begin() as conn:
    conn.execute(
        bridge_agent_sql,
        bridge_agt_df.to_dict(orient="records"),
    )
logging.info(" Agent Bridge table uploaded")

dim_acc_df = df_load[["transaction_key", "Account Number", "Account Name"]]
dim_acc_df["Account Number"] = dim_acc_df["Account Number"].str.split()
dim_acc_df = dim_acc_df.explode("Account Number", ignore_index=True)
dim_acc_df = dim_acc_df.rename(
    columns={
        "Account Number": "account_number",
        "Account Name": "organisation_name",
    }
)
dim_acc_df["account_number"] = dim_acc_df["account_number"].str.strip()
logging.info(" Dim Account Table dataframe created.")

dim_acc_sql = text("""
    INSERT INTO dim_account (
        account_number,
        organisation_name
    )
    VALUES (
        :account_number,
        :organisation_name
    )
    ON CONFLICT (account_number)
    DO UPDATE SET
        account_number = EXCLUDED.account_number
    RETURNING account_key
""")

logging.info(" Creating Account Key column.")
dim_acc_df["account_key"] = pd.Series(pd.NA, index=dim_acc_df.index, dtype="Int64")
logging.info(" Account Key exchange loop initiated.")

with engine.begin() as conn:
    for index, row in dim_acc_df.iterrows():
        account_key = conn.execute(
            dim_acc_sql,
            {
                "account_number": row["account_number"],
                "organisation_name": row["organisation_name"],
            },
        ).scalar_one()

        dim_acc_df.loc[index, "account_key"] = account_key

logging.info(" Account table data uploaded and keys returned")

logging.info(" Creating Account Bridge table.")
bridge_acc_df = dim_acc_df[["transaction_key", "account_key"]]

bridge_acc_sql = text("""
    INSERT INTO bridge_admin_account(
        transaction_key,
        account_key
    )
    VALUES(
        :transaction_key,
        :account_key
    )
    ON CONFLICT (
        transaction_key,
        account_key
    )
    DO NOTHING
""")

logging.info(" Making Account Bridge SQL call.")
with engine.begin() as conn:
    conn.execute(
        bridge_acc_sql,
        bridge_acc_df.to_dict(orient="records"),
    )
logging.info(" Account Bridge table data uploaded.")

logging.info("Creating Audit Extension table and dataframe.")
ext_audit_df = df_load[
    [
        "transaction_key",
        "Audit or Change of Auditor",
        "Description of Request (keep it short)",
        "If 'Returned', why?",
    ]
]

ext_audit_df = ext_audit_df.rename(
    columns={
        "Audit or Change of Auditor": "audit_change_auditor",
        "Description of Request (keep it short)": "dec_req",
        "If 'Returned', why?": "ret_note",
    }
)

ext_audit_sql = text("""
    INSERT INTO ext_fact_admin_audits(
        transaction_key,
        audit_change_auditor,
        dec_req,
        ret_note
    )
    VALUES(
        :transaction_key,
        :audit_change_auditor,
        :dec_req,
        :ret_note
    )
    ON CONFLICT (transaction_key)
    DO UPDATE SET
        audit_change_auditor = EXCLUDED.audit_change_auditor,
        dec_req = EXCLUDED.dec_req,
        ret_note = EXCLUDED.ret_note,
        updated_at = CURRENT_TIMESTAMP
""")

logging.info(" Making Audit Extension table SQL call.")
with engine.begin() as conn:
    conn.execute(
        ext_audit_sql,
        ext_audit_df.to_dict(orient="records"),
    )
logging.info(" Audit Extension table uploaded")